"""DOCX renderer for validated report objects. Does not open external resources."""
from io import BytesIO
from docx import Document
from docx.shared import Pt
from app.services.report_rendering import document_blocks


class DocxReportRenderer:
    media_type = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'

    def render(self, report):
        document = Document()
        document.core_properties.title = report.title
        normal = document.styles['Normal']
        normal.font.name = 'Times New Roman'
        normal.font.size = Pt(12)
        normal.paragraph_format.line_spacing = 2
        for kind, text in document_blocks(report):
            if kind in ('title', 'heading', 'subheading'):
                document.add_heading(text, level={'title': 0, 'heading': 1, 'subheading': 2}[kind])
            else:
                document.add_paragraph(text)
        output = BytesIO()
        document.save(output)
        return output.getvalue()
